# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES.
# All rights reserved.
# SPDX-License-Identifier: Apache-2.0

import io
from typing import Any

import pytest

from auto_ontology.glossary_import import service
from auto_ontology.glossary_import.extractor import (
    ExtractedTerm,
    ExtractionResult,
    extract_terms,
    merge_terms,
)
from auto_ontology.glossary_import.parsers import (
    DocumentError,
    chunk,
    extract_sections,
)


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("销售指标口径说明", level=1)
    document.add_paragraph(
        "GMV：统计周期内已完成订单的实付金额之和，不含已取消和已退货订单。"
    )
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(
        table.rows,
        [
            ("字段名", "中文名", "说明"),
            ("status", "订单状态", "已完成/已取消/已退货"),
            ("channel", "渠道", "线上商城、小程序、线下门店"),
        ],
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据字典"
    sheet.append(["字段名", "中文名", "说明"])
    sheet.append(["member_level", "会员等级", "普通/银卡/金卡/钻石"])
    sheet.append([None, None, None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """A minimal one-page PDF whose content stream draws *text* (ASCII)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


class TestParsers:
    def test_docx_keeps_paragraphs_and_table_headers(self) -> None:
        sections = extract_sections("口径.docx", _docx_bytes())
        text = "\n".join(sections)
        assert "GMV：统计周期内已完成订单的实付金额之和" in text
        assert "字段名: status; 中文名: 订单状态; 说明: 已完成/已取消/已退货" in text
        # The heading comes before the table, in document order.
        assert text.index("销售指标口径说明") < text.index("[表格]")

    def test_xlsx_rows_become_header_value_pairs(self) -> None:
        [section] = extract_sections("dict.xlsx", _xlsx_bytes())
        assert section.startswith("[工作表 数据字典]")
        assert (
            "字段名: member_level; 中文名: 会员等级; 说明: 普通/银卡/金卡/钻石"
            in section
        )

    def test_csv_in_gbk_is_decoded(self) -> None:
        data = "字段名,中文名\norder_id,订单编号\n".encode("gb18030")
        [section] = extract_sections("dict.csv", data)
        assert "字段名: order_id; 中文名: 订单编号" in section

    def test_pdf_text_is_read_per_page(self) -> None:
        [section] = extract_sections(
            "doc.pdf", _pdf_bytes("GMV means gross merchandise value")
        )
        assert section.startswith("[第 1 页]")
        assert "gross merchandise value" in section

    def test_unsupported_and_empty_files_are_refused(self) -> None:
        with pytest.raises(DocumentError, match="不支持"):
            extract_sections("old.doc", b"x")
        with pytest.raises(DocumentError, match="空"):
            extract_sections("a.txt", b"")
        with pytest.raises(DocumentError):
            extract_sections("broken.docx", b"not a zip")


def test_chunk_packs_and_splits() -> None:
    sections = ["a" * 10, "b" * 10, "c" * 25]
    chunks = chunk(sections, max_chars=24)
    assert chunks[0] == "a" * 10 + "\n\n" + "b" * 10
    assert all(len(c) <= 24 for c in chunks)
    assert "".join(c.replace("\n", "") for c in chunks).count("c") == 25


def test_merge_terms_dedupes_and_unions_synonyms() -> None:
    merged = merge_terms(
        [
            [ExtractedTerm(term="GMV", definition="成交", synonyms=["成交总额"])],
            [
                ExtractedTerm(
                    term="gmv ",
                    definition="已完成订单实付金额之和",
                    synonyms=["GMV", "Gross Merchandise Volume"],
                ),
                ExtractedTerm(term="客单价", definition="销售额除以订单数"),
                ExtractedTerm(term="  ", definition="x"),
            ],
        ]
    )
    assert [m["term"] for m in merged] == ["GMV", "客单价"]
    # The fragment "成交" is replaced by the fuller later definition.
    assert merged[0]["definition"] == "已完成订单实付金额之和"
    assert merged[0]["synonyms"] == ["成交总额", "Gross Merchandise Volume"]


class FakeLLM:
    """Stands in for the chat model: answers each chunk from a script."""

    def __init__(self, answers: list[Any]) -> None:
        self.answers = list(answers)


def test_extract_terms_skips_a_failing_chunk(monkeypatch: pytest.MonkeyPatch) -> None:
    answers = [
        ExtractionResult(
            terms=[ExtractedTerm(term="GMV", definition="已完成订单实付金额之和")]
        ),
        RuntimeError("model down"),
        ExtractionResult(
            terms=[ExtractedTerm(term="复购率", definition="下单两次及以上的客户占比")]
        ),
    ]

    def fake_invoke(llm: Any, messages: Any, schema: Any) -> Any:
        answer = answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        return answer

    monkeypatch.setattr(
        "auto_ontology.utils.llm_invoke.invoke_with_structured_output", fake_invoke
    )
    progress: list[int] = []
    merged = extract_terms(["a", "b", "c"], object(), on_chunk=progress.append)
    assert [m["term"] for m in merged] == ["GMV", "复购率"]
    assert progress == [1, 2, 3]


def test_start_import_records_progress_and_candidates(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: dict[str, Any] = {"progress": []}
    monkeypatch.setattr(
        service.dal, "create_batch", lambda **kw: calls.setdefault("batch", kw) and "b1"
    )
    monkeypatch.setattr(
        service.dal, "set_progress", lambda b, n: calls["progress"].append(n)
    )
    monkeypatch.setattr(
        service.dal, "finish_batch", lambda b, c: calls.setdefault("finished", c)
    )
    monkeypatch.setattr(
        "auto_ontology.utils.llm_invoke.invoke_with_structured_output",
        lambda llm, m, s: ExtractionResult(
            terms=[ExtractedTerm(term="订单状态", definition="订单当前所处阶段")]
        ),
    )

    batch_id = service.start_import(
        "口径.docx", _docx_bytes(), "u1", llm_factory=object, background=False
    )

    assert batch_id == "b1"
    assert calls["batch"]["file_name"] == "口径.docx"
    assert calls["batch"]["truncated"] is False
    assert calls["progress"] == [1]
    assert calls["finished"][0]["term"] == "订单状态"


def test_start_import_rejects_oversized_files() -> None:
    with pytest.raises(DocumentError, match="20 MB"):
        service.start_import(
            "big.txt", b"x" * (service.MAX_BYTES + 1), None, llm_factory=object
        )
